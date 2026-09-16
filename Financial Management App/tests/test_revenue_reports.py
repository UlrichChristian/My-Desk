"""Unit tests for Biggest Clients rollup and income exclusions."""

from __future__ import annotations

import pandas as pd

from fin_services.account_integrations import create_integration
from fin_services.revenue_reports import (
    EXCLUDED_INCOME_ACCOUNTS,
    biggest_clients,
    client_period_detail,
    data_from_db,
    filter_by_client_search,
    filter_by_integration,
    report_from_db,
    to_excel,
)
from tests.helpers import insert_batch, insert_revenue


def test_excluded_income_accounts():
    assert "Referral Fees:Email Feed" in EXCLUDED_INCOME_ACCOUNTS
    assert "Statement of Work" in EXCLUDED_INCOME_ACCOUNTS


def test_client_period_detail_by_income_bucket(conn):
    batch = insert_batch(conn)
    insert_revenue(
        conn, period="2026-03", client_key="Acme", amount=100.0,
        quantity=10.0, rate=10.0, product_code="ADMIN",
        income_account="Admin Fees (earned on Premium)",
        premium_split=50.0, lives_split=5.0, import_id=batch,
    )
    insert_revenue(
        conn, period="2026-03", client_key="Acme", amount=40.0,
        quantity=8.0, rate=5.0, product_code="SUB",
        income_account="Admin Fees (earned on Sub. Fees)",
        premium_split=0.0, lives_split=0.0, import_id=batch,
    )
    insert_revenue(
        conn, period="2026-03", client_key="Acme", amount=20.0,
        quantity=4.0, rate=5.0, product_code="SUB-B",
        income_account="Admin Fees (earned on Sub. Fees)",
        premium_split=0.0, lives_split=0.0, import_id=batch,
    )
    insert_revenue(
        conn, period="2026-03", client_key="Acme", amount=999.0,
        income_account="Statement of Work", import_id=batch,
    )
    detail = client_period_detail(conn, "2026-03", "Acme")
    assert detail["client"] == "Acme"
    assert detail["period"] == "2026-03"
    accounts = {b["income_account"] for b in detail["buckets"]}
    assert "Statement of Work" not in accounts
    assert len(detail["buckets"]) == 3
    products = {(b["income_account"], b["product_code"]) for b in detail["buckets"]}
    assert ("Admin Fees (earned on Sub. Fees)", "SUB") in products
    assert ("Admin Fees (earned on Sub. Fees)", "SUB-B") in products
    by_prod = {(b["income_account"], b["product_code"]): b for b in detail["buckets"]}
    assert by_prod[("Admin Fees (earned on Premium)", "ADMIN")]["certs"] == 10.0
    assert by_prod[("Admin Fees (earned on Sub. Fees)", "SUB")]["amount"] == 40.0


def test_biggest_clients_ranks_and_shares():
    enriched = pd.DataFrame([
        {
            "period": "2026-03",
            "client_key": "Alpha",
            "advisor": "A",
            "amount": 300.0,
            "premium_split": 100.0,
            "lives_split": 10.0,
            "income_account": "Admin Fees",
        },
        {
            "period": "2026-03",
            "client_key": "Beta",
            "advisor": "B",
            "amount": 100.0,
            "premium_split": 50.0,
            "lives_split": 5.0,
            "income_account": "Admin Fees",
        },
    ])
    rep = biggest_clients(enriched, "2026-03")
    assert list(rep["Client"]) == ["Alpha", "Beta"]
    assert rep["Revenue"].tolist() == [300.0, 100.0]
    assert abs(rep["Revenue Share"].sum() - 1.0) < 1e-9
    assert rep.attrs["total_revenue"] == 400.0


def test_biggest_clients_excludes_non_client_income():
    enriched = pd.DataFrame([
        {
            "period": "2026-03",
            "client_key": "Manulife",
            "advisor": None,
            "amount": 500.0,
            "premium_split": 0.0,
            "lives_split": 0.0,
            "income_account": "Referral Fees:Email Feed",
        },
        {
            "period": "2026-03",
            "client_key": "SOW Client",
            "advisor": None,
            "amount": 1000.0,
            "premium_split": 0.0,
            "lives_split": 0.0,
            "income_account": "Statement of Work",
        },
        {
            "period": "2026-03",
            "client_key": "Real Client",
            "advisor": "A",
            "amount": 200.0,
            "premium_split": 100.0,
            "lives_split": 10.0,
            "income_account": "Admin Fees earned on Premium",
        },
    ])
    rep = biggest_clients(enriched, "2026-03")
    assert list(rep["Client"]) == ["Real Client"]
    assert rep.attrs["total_revenue"] == 200.0


def test_report_from_db_excludes_referral_fees(conn):
    batch = insert_batch(conn)
    insert_revenue(
        conn,
        period="2026-03",
        client_key="Manulife",
        amount=999.0,
        income_account="Referral Fees:Email Feed",
        import_id=batch,
    )
    insert_revenue(
        conn,
        period="2026-03",
        client_key="SOW Co",
        amount=888.0,
        income_account="Statement of Work",
        import_id=batch,
    )
    insert_revenue(
        conn,
        period="2026-03",
        client_key="Acme",
        amount=100.0,
        import_id=batch,
        premium_split=50.0,
        lives_split=5.0,
    )
    rep = report_from_db(conn, "2026-03")
    assert list(rep["Client"]) == ["Acme"]
    assert rep.attrs["total_revenue"] == 100.0


def test_report_from_db_empty(conn):
    rep = report_from_db(conn)
    assert rep.empty
    assert rep.attrs["period"] is None


def test_report_from_db_integration_flags_and_filters(conn):
    batch = insert_batch(conn)
    insert_revenue(
        conn, period="2026-03", client_key="Acme", amount=100.0,
        oid="1", import_id=batch, premium_split=50.0, lives_split=5.0,
    )
    insert_revenue(
        conn, period="2026-03", client_key="Beta", amount=80.0,
        oid="2", import_id=batch,
    )
    insert_revenue(
        conn, period="2026-03", client_key="Gamma", amount=60.0,
        oid="3", import_id=batch,
    )
    create_integration(conn, client="Acme", vendor_name="ADP", role="payroll", status="stable")
    create_integration(
        conn, client="Beta", vendor_name="Raw HRIS", role="hris", status="on_trial",
    )
    create_integration(
        conn, client="Gamma", vendor_name="ADP", role="payroll", status="in_discussion",
    )

    rep = report_from_db(conn, "2026-03")
    by_client = {row["Client"]: row for _, row in rep.iterrows()}
    assert bool(by_client["Acme"]["Payroll"]) is True
    assert bool(by_client["Acme"]["HRIS"]) is False
    assert bool(by_client["Beta"]["Payroll"]) is False
    assert bool(by_client["Beta"]["HRIS"]) is True
    assert bool(by_client["Gamma"]["Payroll"]) is False
    assert bool(by_client["Gamma"]["HRIS"]) is False

    assert set(filter_by_integration(rep, "none")["Client"]) == {"Gamma"}
    assert set(filter_by_integration(rep, "no_payroll")["Client"]) == {"Beta", "Gamma"}
    assert set(filter_by_integration(rep, "no_hris")["Client"]) == {"Acme", "Gamma"}
    assert set(filter_by_integration(rep, "all")["Client"]) == {"Acme", "Beta", "Gamma"}
    assert set(filter_by_client_search(rep, "ac")["Client"]) == {"Acme"}
    assert set(filter_by_client_search(rep, "")["Client"]) == {"Acme", "Beta", "Gamma"}


def test_excel_client_sheet_writes_yes_no_flags(conn, tmp_path):
    batch = insert_batch(conn)
    insert_revenue(
        conn, period="2026-03", client_key="Acme", amount=100.0,
        oid="1", import_id=batch,
    )
    create_integration(conn, client="Acme", vendor_name="ADP", role="payroll", status="stable")
    rep = report_from_db(conn, "2026-03")
    path = tmp_path / "clients.xlsx"
    to_excel(rep, path, period="2026-03", native_pivots=False)

    from openpyxl import load_workbook
    ws = load_workbook(path)["Revenue by Client"]
    headers = [ws.cell(7, col).value for col in range(2, 12)]
    assert headers[:4] == ["Client", "Advisor", "Payroll", "HRIS"]
    assert ws.cell(8, 2).value == "Acme"
    assert ws.cell(8, 4).value == "Yes"
    assert ws.cell(8, 5).value == "No"

    data = data_from_db(conn, "2026-03")
    assert set(data["Payroll"]) == {"Yes"}
    assert set(data["HRIS"]) == {"No"}
    data_path = tmp_path / "with_data.xlsx"
    to_excel(rep, data_path, period="2026-03", data=data, native_pivots=False)
    data_headers = [c.value for c in load_workbook(data_path)["Data"][1]]
    assert "Payroll" in data_headers
    assert "HRIS" in data_headers
    pay_col = data_headers.index("Payroll") + 1
    hris_col = data_headers.index("HRIS") + 1
    assert load_workbook(data_path)["Data"].cell(2, pay_col).value == "Yes"
    assert load_workbook(data_path)["Data"].cell(2, hris_col).value == "No"


def test_filter_by_client_search_matches_substring(conn):
    batch = insert_batch(conn)
    insert_revenue(
        conn, period="2026-03", client_key="ADP Services", amount=10.0,
        oid="1", import_id=batch,
    )
    insert_revenue(
        conn, period="2026-03", client_key="New Gold", amount=20.0,
        oid="2", import_id=batch,
    )
    rep = report_from_db(conn, "2026-03")
    matched = filter_by_client_search(rep, "dp")
    assert list(matched["Client"]) == ["ADP Services"]
    assert matched.attrs["total_revenue"] == 10.0
    assert list(filter_by_client_search(rep, "DP")["Client"]) == ["ADP Services"]


def test_report_from_db_blanks_premium_when_metrics_are_missing(conn):
    """A month with revenue but no vintage must not show today's premium as $0."""
    batch = insert_batch(conn)
    insert_revenue(
        conn, period="2025-08", client_key="Acme", amount=100.0, import_id=batch,
    )
    rep = report_from_db(conn, "2025-08")
    assert list(rep["Client"]) == ["Acme"]
    assert rep.attrs["total_revenue"] == 100.0
    assert pd.isna(rep["Premium*"].iloc[0])
    assert pd.isna(rep["Lives*"].iloc[0])
    assert pd.isna(rep["Revenue/Premium"].iloc[0])
    assert pd.isna(rep["Revenue/Life"].iloc[0])


def test_client_period_detail_blanks_premium_when_metrics_are_missing(conn):
    batch = insert_batch(conn)
    insert_revenue(
        conn, period="2025-08", client_key="Acme", amount=100.0,
        quantity=10.0, rate=10.0, product_code="ADMIN",
        income_account="Admin Fees (earned on Premium)", import_id=batch,
    )
    detail = client_period_detail(conn, "2025-08", "Acme")
    assert len(detail["buckets"]) == 1
    assert detail["buckets"][0]["premium"] is None
    assert detail["buckets"][0]["lives"] is None
