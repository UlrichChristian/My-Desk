"""Month-over-month revenue change bridge (new/lost + certs/price/benefits).

Flask-free: takes a DB connection and returns summary + client-level DataFrame
so the web UI, Excel export, and a future MCP tool can share the same engine.

Driver math for a matched (client_key, product_code) pair with usable qty/rate:
  Certs (volume) = (q1 - q0) * r0
  Price          = q1 * (r1 - r0)
  (Note: Certs + Price = q1*r1 - q0*r0 when amount ≈ qty*rate.)

Unmatched product lines on an existing client → Benefits (mix).
Rows missing qty/rate → Other (so the bridge still ties to Δ revenue).
"""

from __future__ import annotations

from calendar import month_name
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from fin_services.revenue_reports import EXCLUDED_INCOME_ACCOUNTS

CLIENT_COLUMNS = [
    "Client", "Advisor",
    "Current Revenue", "Prior Revenue", "Δ Revenue",
    "New", "Lost", "Certs", "Price", "Benefits", "Other",
]

DRIVER_KEYS = ("new", "lost", "certs", "price", "benefits", "other")

_FONT = "Lato"
_HEADER_FONT = Font(name=_FONT, size=9, bold=True, color="FFFFFF")
_CELL_FONT = Font(name=_FONT, size=9)
_TITLE_FONT = Font(name=_FONT, size=9, bold=True)
_HEADER_FILL = PatternFill("solid", fgColor="5B9BD5")
_THIN = Border(
    left=Side(style="thin", color="B4B4B4"),
    right=Side(style="thin", color="B4B4B4"),
    top=Side(style="thin", color="B4B4B4"),
    bottom=Side(style="thin", color="B4B4B4"),
)


def _first_nonnull(series: pd.Series):
    s = series.dropna()
    return s.iloc[0] if len(s) else None


def _prior_period(period: str) -> str | None:
    """Return YYYY-MM one calendar month before ``period``, or None if unparsable."""
    try:
        y, m = int(period[:4]), int(period[5:7])
        if m == 1:
            return f"{y - 1:04d}-12"
        return f"{y:04d}-{m - 1:02d}"
    except (TypeError, ValueError, IndexError):
        return None


def _period_label(period: str | None) -> str:
    if not period or len(period) < 7:
        return period or ""
    try:
        return f"{month_name[int(period[5:7])]} {period[:4]}"
    except (TypeError, ValueError, IndexError):
        return period


def available_revenue_periods(conn) -> list[str]:
    return [
        r[0] for r in conn.execute(
            "SELECT DISTINCT period FROM fact_revenue WHERE period IS NOT NULL ORDER BY period DESC"
        )
    ]


def available_change_periods(conn) -> list[str]:
    """Periods usable as the current month (need ≥1 other period as a reference)."""
    periods = available_revenue_periods(conn)
    return periods if len(periods) >= 2 else []


def default_prior_period(period: str, present) -> str | None:
    """Calendar prior month when available; else latest earlier period; else any other."""
    present_set = set(present)
    cal = _prior_period(period)
    if cal and cal in present_set:
        return cal
    earlier = sorted(p for p in present_set if p < period)
    if earlier:
        return earlier[-1]
    others = sorted(p for p in present_set if p != period)
    return others[0] if others else None


def resolve_prior_period(period: str, prior_period: str | None, present) -> str | None:
    """Validate an explicit prior, otherwise fall back to ``default_prior_period``."""
    present_set = set(present)
    if prior_period and prior_period in present_set and prior_period != period:
        return prior_period
    return default_prior_period(period, present_set)


def _aggregate_period(conn, period: str) -> pd.DataFrame:
    """One row per (client_key, product_code) for a period with amount/qty/rate."""
    placeholders = ",".join("?" * len(EXCLUDED_INCOME_ACCOUNTS))
    rows = conn.execute(
        "SELECT client_key, product_code, advisor_key, amount, quantity, rate "
        f"FROM fact_revenue WHERE period = ? AND COALESCE(income_account, '') NOT IN ({placeholders})",
        (period, *sorted(EXCLUDED_INCOME_ACCOUNTS)),
    ).fetchall()
    raw = pd.DataFrame([dict(r) for r in rows])
    if raw.empty:
        return pd.DataFrame(columns=[
            "client_key", "product_code", "advisor", "amount", "quantity", "rate",
        ])

    raw["product_code"] = raw["product_code"].fillna("(blank)")
    grp = raw.groupby(["client_key", "product_code"], dropna=False)
    amt = grp["amount"].sum()
    qty = grp["quantity"].sum(min_count=1)
    # Effective rate from totals when qty is usable; else leave null.
    rate = amt / qty.replace(0, pd.NA)
    out = pd.DataFrame({
        "client_key": amt.index.get_level_values(0),
        "product_code": amt.index.get_level_values(1),
        "advisor": grp["advisor_key"].agg(_first_nonnull).values,
        "amount": amt.values,
        "quantity": qty.values,
        "rate": rate.values,
    })
    return out.reset_index(drop=True)


def _usable_qr(q, r) -> bool:
    return pd.notna(q) and pd.notna(r) and float(q) != 0


def _is_missing_client(c) -> bool:
    return c is None or (isinstance(c, float) and pd.isna(c))


def _client_mask(series: pd.Series, client) -> pd.Series:
    if _is_missing_client(client):
        return series.isna()
    return series == client


def _bridge_existing(prior_prod: pd.DataFrame, curr_prod: pd.DataFrame) -> dict[str, float]:
    """Split Δ for an existing client across Certs / Price / Benefits / Other."""
    certs = price = benefits = other = 0.0
    prior_idx = prior_prod.set_index("product_code")
    curr_idx = curr_prod.set_index("product_code")
    all_products = set(prior_idx.index) | set(curr_idx.index)

    for prod in all_products:
        p = prior_idx.loc[prod] if prod in prior_idx.index else None
        c = curr_idx.loc[prod] if prod in curr_idx.index else None
        # groupby index can yield Series or (rarely) DataFrame if duplicate keys —
        # aggregation guarantees uniqueness, but guard anyway.
        if isinstance(p, pd.DataFrame):
            p = p.iloc[0]
        if isinstance(c, pd.DataFrame):
            c = c.iloc[0]

        if p is None:
            # New product line for existing client → Benefits
            benefits += float(c["amount"])
            continue
        if c is None:
            # Dropped product line → Benefits (negative)
            benefits -= float(p["amount"])
            continue

        a0, a1 = float(p["amount"]), float(c["amount"])
        q0, q1 = p["quantity"], c["quantity"]
        r0, r1 = p["rate"], c["rate"]

        if _usable_qr(q0, r0) and _usable_qr(q1, r1):
            q0f, q1f, r0f, r1f = float(q0), float(q1), float(r0), float(r1)
            vol = (q1f - q0f) * r0f
            prc = q1f * (r1f - r0f)
            certs += vol
            price += prc
            # Residual from amount ≠ qty*rate (credits, rounding) → Other
            explained = vol + prc
            residual = (a1 - a0) - explained
            if abs(residual) > 0.005:
                other += residual
        else:
            other += a1 - a0

    return {"certs": certs, "price": price, "benefits": benefits, "other": other}


def change_report(
    conn,
    period: str | None = None,
    prior_period: str | None = None,
) -> dict:
    """Build the MoM change report for ``period`` vs ``prior_period``.

    ``prior_period`` defaults to the calendar month before ``period`` when that
    month is in the store; otherwise the latest earlier available month.

    Returns dict with keys:
      period, prior_period, summary (driver totals), clients (DataFrame)
    """
    all_periods = available_revenue_periods(conn)
    periods = available_change_periods(conn)
    if not periods:
        empty = pd.DataFrame(columns=CLIENT_COLUMNS)
        return {
            "period": None,
            "prior_period": None,
            "summary": _empty_summary(),
            "clients": empty,
        }

    if period is None or period not in all_periods:
        period = periods[0]
    prior = resolve_prior_period(period, prior_period, all_periods)
    if prior is None:
        empty = pd.DataFrame(columns=CLIENT_COLUMNS)
        return {
            "period": period,
            "prior_period": None,
            "summary": _empty_summary(),
            "clients": empty,
        }

    prior_df = _aggregate_period(conn, prior)
    curr_df = _aggregate_period(conn, period)

    prior_clients = set(prior_df["client_key"].unique()) if not prior_df.empty else set()
    curr_clients = set(curr_df["client_key"].unique()) if not curr_df.empty else set()

    def _client_sort_key(c):
        if c is None or (isinstance(c, float) and pd.isna(c)):
            return (1, "")
        return (0, str(c))

    all_clients = sorted(prior_clients | curr_clients, key=_client_sort_key)

    rows = []
    for client in all_clients:
        p_slice = prior_df[_client_mask(prior_df["client_key"], client)] if not prior_df.empty else prior_df
        c_slice = curr_df[_client_mask(curr_df["client_key"], client)] if not curr_df.empty else curr_df
        prior_rev = float(p_slice["amount"].sum()) if len(p_slice) else 0.0
        curr_rev = float(c_slice["amount"].sum()) if len(c_slice) else 0.0
        advisor = _first_nonnull(c_slice["advisor"]) if len(c_slice) else None
        if advisor is None and len(p_slice):
            advisor = _first_nonnull(p_slice["advisor"])

        new = lost = certs = price = benefits = other = 0.0
        # Membership via set: NaN != NaN, so handle missing explicitly.
        if _is_missing_client(client):
            in_prior = any(_is_missing_client(c) for c in prior_clients)
            in_curr = any(_is_missing_client(c) for c in curr_clients)
        else:
            in_prior = client in prior_clients
            in_curr = client in curr_clients

        if in_curr and not in_prior:
            new = curr_rev
        elif in_prior and not in_curr:
            lost = -prior_rev
        else:
            drivers = _bridge_existing(p_slice, c_slice)
            certs = drivers["certs"]
            price = drivers["price"]
            benefits = drivers["benefits"]
            other = drivers["other"]

        rows.append({
            "Client": "(unmapped)" if _is_missing_client(client) else client,
            "Advisor": advisor or "",
            "Current Revenue": curr_rev,
            "Prior Revenue": prior_rev,
            "Δ Revenue": curr_rev - prior_rev,
            "New": new,
            "Lost": lost,
            "Certs": certs,
            "Price": price,
            "Benefits": benefits,
            "Other": other,
        })

    clients = pd.DataFrame(rows, columns=CLIENT_COLUMNS)
    if not clients.empty:
        clients = clients.iloc[
            clients["Δ Revenue"].abs().sort_values(ascending=False).index
        ].reset_index(drop=True)

    prior_total = float(clients["Prior Revenue"].sum()) if len(clients) else 0.0
    curr_total = float(clients["Current Revenue"].sum()) if len(clients) else 0.0
    delta = curr_total - prior_total
    drivers_sum = {
        k: float(clients[col].sum()) if len(clients) else 0.0
        for k, col in zip(DRIVER_KEYS, ("New", "Lost", "Certs", "Price", "Benefits", "Other"))
    }
    explained = sum(drivers_sum.values())
    summary = {
        "prior_revenue": prior_total,
        "current_revenue": curr_total,
        "delta": delta,
        **drivers_sum,
        "explained": explained,
        "residual": delta - explained,
        "period_label": _period_label(period),
        "prior_label": _period_label(prior),
    }
    return {
        "period": period,
        "prior_period": prior,
        "summary": summary,
        "clients": clients,
    }


def _empty_summary() -> dict:
    return {
        "prior_revenue": 0.0,
        "current_revenue": 0.0,
        "delta": 0.0,
        "new": 0.0,
        "lost": 0.0,
        "certs": 0.0,
        "price": 0.0,
        "benefits": 0.0,
        "other": 0.0,
        "explained": 0.0,
        "residual": 0.0,
        "period_label": "",
        "prior_label": "",
    }


UNMAPPED_CLIENT_LABEL = "(unmapped)"

_DETAIL_COLS = (
    "period", "txn_date", "txn_type", "invoice_no", "name_raw", "memo",
    "product_code", "income_account", "quantity", "rate", "amount",
)


def _is_unmapped_client(client_key) -> bool:
    return (
        client_key is None
        or (isinstance(client_key, float) and pd.isna(client_key))
        or str(client_key).strip() == UNMAPPED_CLIENT_LABEL
    )


def _fetch_client_lines(conn, periods: list[str], client_key: str) -> list[dict]:
    """Raw fact_revenue lines for a client in the given periods (exclusions applied)."""
    if not periods:
        return []
    placeholders = ",".join("?" * len(EXCLUDED_INCOME_ACCOUNTS))
    period_ph = ",".join("?" * len(periods))
    exclude = tuple(sorted(EXCLUDED_INCOME_ACCOUNTS))
    unmapped = _is_unmapped_client(client_key)

    if unmapped:
        sql = (
            "SELECT period, txn_date, txn_type, invoice_no, name_raw, memo, "
            "product_code, income_account, quantity, rate, amount "
            "FROM fact_revenue "
            f"WHERE period IN ({period_ph}) AND client_key IS NULL "
            f"AND COALESCE(income_account, '') NOT IN ({placeholders}) "
            "ORDER BY period ASC, income_account ASC, product_code ASC, txn_date ASC"
        )
        params = (*periods, *exclude)
    else:
        sql = (
            "SELECT period, txn_date, txn_type, invoice_no, name_raw, memo, "
            "product_code, income_account, quantity, rate, amount "
            "FROM fact_revenue "
            f"WHERE period IN ({period_ph}) AND client_key = ? "
            f"AND COALESCE(income_account, '') NOT IN ({placeholders}) "
            "ORDER BY period ASC, income_account ASC, product_code ASC, txn_date ASC"
        )
        params = (*periods, client_key, *exclude)

    rows = []
    for r in conn.execute(sql, params).fetchall():
        d = dict(r) if not isinstance(r, dict) else r
        rows.append({
            col: (None if d.get(col) is None or (isinstance(d.get(col), float) and pd.isna(d.get(col)))
                  else d.get(col))
            for col in _DETAIL_COLS
        })
    return rows


def _period_metrics(lines: list[dict]) -> dict:
    """Sum amount/quantity for a set of lines; effective rate = amount/qty."""
    if not lines:
        return {"quantity": None, "rate": None, "amount": 0.0, "certs": None}
    amt = sum(float(x["amount"] or 0) for x in lines)
    qtys = [float(x["quantity"]) for x in lines if x.get("quantity") is not None]
    qty = sum(qtys) if qtys else None
    rate = None
    if qty is not None and qty != 0:
        rate = amt / qty
    # Display certs as absolute headcount (QBO often stores qty negative).
    certs = abs(qty) if qty is not None else None
    return {"quantity": qty, "rate": rate, "amount": amt, "certs": certs}


def _compare_buckets(lines: list[dict], prior: str, current: str) -> list[dict]:
    """One row per (income_account, product_code) with prior/current metrics."""
    by_key: dict[tuple[str, str], list[dict]] = {}
    for line in lines:
        income = line.get("income_account") or "(blank)"
        product = line.get("product_code") or "(blank)"
        by_key.setdefault((income, product), []).append(line)

    buckets = []
    for (income_account, product_code), bucket_lines in by_key.items():
        prior_m = _period_metrics([x for x in bucket_lines if x.get("period") == prior])
        curr_m = _period_metrics([x for x in bucket_lines if x.get("period") == current])
        if prior_m["amount"] == 0 and curr_m["amount"] == 0 and prior_m["certs"] is None and curr_m["certs"] is None:
            continue
        buckets.append({
            "income_account": income_account,
            "product_code": product_code,
            "prior": prior_m,
            "current": curr_m,
            "delta_amount": curr_m["amount"] - prior_m["amount"],
        })
    # Group visually by income type, largest movers first within each type.
    buckets.sort(key=lambda b: (b["income_account"], -abs(b["delta_amount"]), b["product_code"]))
    return buckets


def client_detail_rows(
    conn,
    period: str,
    client_key: str,
    prior_period: str | None = None,
) -> dict:
    """Side-by-side product compare for a client across two months.

    Returns ``buckets`` (one per income type + product/service with prior/current
    certs, rate, amount) plus raw ``rows`` for audit. ``client_key`` may be
    ``(unmapped)``.
    """
    present = available_revenue_periods(conn)
    prior = resolve_prior_period(period, prior_period, present)
    display_client = UNMAPPED_CLIENT_LABEL if _is_unmapped_client(client_key) else client_key
    if prior is None:
        return {
            "period": period,
            "prior_period": None,
            "client": display_client,
            "buckets": [],
            "rows": [],
        }

    rows = _fetch_client_lines(conn, [prior, period], client_key)
    return {
        "period": period,
        "prior_period": prior,
        "client": display_client,
        "buckets": _compare_buckets(rows, prior, period),
        "rows": rows,
    }


def to_excel(report: dict, path: str | Path) -> Path:
    """Write summary + client detail sheets for the change report."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    clients = report["clients"]
    summary = report["summary"]
    period = report["period"]
    prior = report["prior_period"]

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws["B2"] = f"{summary.get('period_label') or period or 'Change'} vs {summary.get('prior_label') or prior or 'prior'}"
    ws["B2"].font = _TITLE_FONT
    ws["B3"] = "Monthly Change Report"
    ws["B3"].font = _TITLE_FONT

    labels = [
        ("Current Revenue", summary["current_revenue"]),
        ("Prior Revenue", summary["prior_revenue"]),
        ("Δ Revenue", summary["delta"]),
        ("New clients", summary["new"]),
        ("Lost clients", summary["lost"]),
        ("Certs", summary["certs"]),
        ("Price (admin fee)", summary["price"]),
        ("Benefits", summary["benefits"]),
        ("Other", summary["other"]),
        ("Explained", summary["explained"]),
        ("Residual", summary["residual"]),
    ]
    ws["B5"] = "Metric"
    ws["C5"] = "Amount"
    for cell in (ws["B5"], ws["C5"]):
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.border = _THIN
    for i, (label, val) in enumerate(labels, 6):
        ws.cell(i, 2, label).font = _CELL_FONT
        ws.cell(i, 2).border = _THIN
        c = ws.cell(i, 3, val)
        c.font = _CELL_FONT
        c.number_format = "#,##0.00"
        c.border = _THIN
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 14

    ws_c = wb.create_sheet("By Client")
    headers = list(CLIENT_COLUMNS)
    for i, h in enumerate(headers, 1):
        cell = ws_c.cell(1, i, h)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.border = _THIN
        cell.alignment = Alignment(wrap_text=True, horizontal="center")
    money_cols = set(range(3, 12))  # C..K
    for r_i, row in enumerate(clients.itertuples(index=False), 2):
        for c_i, val in enumerate(row, 1):
            cell = ws_c.cell(r_i, c_i, None if pd.isna(val) else val)
            cell.font = _CELL_FONT
            cell.border = _THIN
            if c_i in money_cols:
                cell.number_format = "#,##0.00"
    widths = [22, 18, 12, 12, 11, 10, 10, 10, 10, 10, 10]
    for i, w in enumerate(widths, 1):
        ws_c.column_dimensions[get_column_letter(i)].width = w

    wb.save(path)
    return path
